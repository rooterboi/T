import asyncio
import uuid
from pathlib import Path

from pptx import Presentation
from pptx.util import Pt

from config import DATA_DIR


def _build(outline: dict) -> Path:
    prs = Presentation()
    title = prs.slides.add_slide(prs.slide_layouts[0])
    title.shapes.title.text = outline.get("title", "Taqdimot")
    title.placeholders[1].text = "AI yordamida tayyorlandi"
    for s in outline.get("slides", []):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = s.get("title", "")
        tf = slide.placeholders[1].text_frame
        tf.clear()
        for i, bullet in enumerate(s.get("bullets", [])):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = str(bullet)
            p.font.size = Pt(22)
    path = DATA_DIR / "out" / f"slayd_{uuid.uuid4().hex[:8]}.pptx"
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    return path


async def build_pptx(outline: dict) -> Path:
    return await asyncio.to_thread(_build, outline)
